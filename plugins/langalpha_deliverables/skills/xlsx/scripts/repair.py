#!/usr/bin/env python3
# DALTON-HOOK(skills-registry): repair corrupted definedNames in uploaded xlsx workbooks
"""Repair an uploaded .xlsx workbook whose corrupt definedNames fail openpyxl loading.

Usage:
    python repair.py <file.xlsx> [--out <output.xlsx>]

First attempts `openpyxl.load_workbook(<file.xlsx>)`. If the workbook already loads,
leaves everything untouched, prints "no repair needed", and exits 0.
Otherwise copies `<file.xlsx>` to `<stem>.repaired.xlsx` (or `--out`) via `zipfile`,
strips only `<definedName>` entries that break openpyxl loading (`_xlnm.Print_Titles` /
`_xlnm.Print_Area` rejected by `PrintTitles.from_string` / `PrintArea.from_string`, or
entries with an out-of-range `localSheetId`), verifies that the repaired workbook loads
in openpyxl, and prints the removed names.
"""

from __future__ import annotations

import argparse
import json
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

import openpyxl
from openpyxl.worksheet.print_settings import PrintArea, PrintTitles

try:
    from lxml import etree as _lxml_etree

    _HAS_LXML = True
except Exception:
    _HAS_LXML = False

SML_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
WORKBOOK_XML_PATH = "xl/workbook.xml"


def default_output_path(input_path: Path) -> Path:
    """Return `<stem>.repaired.xlsx` alongside `input_path`."""
    suffix = input_path.suffix if input_path.suffix.lower() in {".xlsx", ".xlsm"} else ".xlsx"
    return input_path.with_name(f"{input_path.stem}.repaired{suffix}")


def _check_defined_name_invalid(
    name: str | None,
    value: str | None,
    local_sheet_id: str | None,
    sheet_count: int,
) -> str | None:
    """Return the reason a `<definedName>` prevents openpyxl loading, or None to keep it."""
    if local_sheet_id is not None:
        try:
            parsed_sheet_id = int(local_sheet_id)
        except ValueError:
            return "invalid_local_sheet_id"
        if parsed_sheet_id < 0 or (sheet_count > 0 and parsed_sheet_id >= sheet_count):
            return "out_of_range_local_sheet_id"

    raw = (value or "").strip()
    if name == "_xlnm.Print_Titles":
        try:
            PrintTitles.from_string(raw)
        except Exception as exc:
            return f"invalid_print_titles: {exc}"

    elif name == "_xlnm.Print_Area":
        try:
            PrintArea.from_string(raw)
        except Exception as exc:
            return f"invalid_print_area: {exc}"

    return None


def _strip_invalid_defined_names_xml(xml_bytes: bytes) -> tuple[bytes, list[dict[str, str]]]:
    """Remove `<definedName>` nodes that fail openpyxl loading from `xl/workbook.xml`."""
    removed: list[dict[str, str]] = []

    if _HAS_LXML:
        parser = _lxml_etree.XMLParser(remove_blank_text=False, resolve_entities=False)
        root = _lxml_etree.fromstring(xml_bytes, parser=parser)
        sheets_tag = f"{{{SML_NS}}}sheet"
        sheet_count = len(root.findall(f".//{sheets_tag}"))
        defined_names_tag = f"{{{SML_NS}}}definedNames"
        defined_name_tag = f"{{{SML_NS}}}definedName"

        for dn_container in list(root.findall(f".//{defined_names_tag}")):
            for dn in list(dn_container.findall(defined_name_tag)):
                name = dn.get("name")
                local_sheet_id = dn.get("localSheetId")
                value = dn.text or ""
                reason = _check_defined_name_invalid(name, value, local_sheet_id, sheet_count)
                if reason is not None:
                    removed.append(
                        {
                            "name": name or "<unnamed>",
                            "value": value,
                            "reason": reason,
                        }
                    )
                    dn_container.remove(dn)
            if len(dn_container) == 0:
                parent = dn_container.getparent()
                if parent is not None:
                    parent.remove(dn_container)

        if not removed:
            return xml_bytes, removed
        out_bytes = _lxml_etree.tostring(
            root,
            xml_declaration=True,
            encoding="UTF-8",
            standalone="yes",
        )
        return out_bytes, removed

    ET.register_namespace("", SML_NS)
    ET.register_namespace("r", "http://schemas.openxmlformats.org/officeDocument/2006/relationships")
    root_et = ET.fromstring(xml_bytes)
    sheets_tag = f"{{{SML_NS}}}sheet"
    sheet_count = len(root_et.findall(f".//{sheets_tag}"))
    defined_names_tag = f"{{{SML_NS}}}definedNames"
    defined_name_tag = f"{{{SML_NS}}}definedName"

    for dn_container in list(root_et.findall(f".//{defined_names_tag}")):
        for dn in list(dn_container.findall(defined_name_tag)):
            name = dn.get("name")
            local_sheet_id = dn.get("localSheetId")
            value = dn.text or ""
            reason = _check_defined_name_invalid(name, value, local_sheet_id, sheet_count)
            if reason is not None:
                removed.append(
                    {
                        "name": name or "<unnamed>",
                        "value": value,
                        "reason": reason,
                    }
                )
                dn_container.remove(dn)
        if len(dn_container) == 0:
            root_et.remove(dn_container)

    if not removed:
        return xml_bytes, removed
    out_bytes = ET.tostring(root_et, encoding="utf-8", xml_declaration=True)
    return out_bytes, removed


def repair_workbook(
    input_path: str | Path,
    output_path: str | Path | None = None,
) -> dict[str, object]:
    """Repair `input_path` into `<stem>.repaired.xlsx` if openpyxl cannot load it directly."""
    src = Path(input_path).resolve()
    if not src.is_file():
        raise FileNotFoundError(f"Workbook not found: {src}")

    try:
        wb = openpyxl.load_workbook(src)
        wb.close()
        return {
            "status": "no_repair_needed",
            "message": "no repair needed",
            "input": str(src),
            "output": None,
            "removed_count": 0,
            "removed_names": [],
            "removed": [],
        }
    except Exception as load_err:
        initial_error = f"{type(load_err).__name__}: {load_err}"

    dst = Path(output_path).resolve() if output_path is not None else default_output_path(src)
    if dst == src:
        raise ValueError("Output path must differ from input path; original file must remain untouched.")

    dst.parent.mkdir(parents=True, exist_ok=True)

    with zipfile.ZipFile(src, "r") as zin:
        namelist = zin.namelist()
        if WORKBOOK_XML_PATH not in namelist:
            raise RuntimeError(
                f"Workbook failed to load ({initial_error}) and has no {WORKBOOK_XML_PATH} to repair."
            )

        workbook_xml = zin.read(WORKBOOK_XML_PATH)
        repaired_xml, removed = _strip_invalid_defined_names_xml(workbook_xml)

        with zipfile.ZipFile(dst, "w", compression=zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = repaired_xml if item.filename == WORKBOOK_XML_PATH else zin.read(item.filename)
                zout.writestr(item, data)

    try:
        repaired_wb = openpyxl.load_workbook(dst)
        repaired_wb.close()
    except Exception as post_err:
        remaining = f"{type(post_err).__name__}: {post_err}"
        if getattr(post_err, "__cause__", None) is not None:
            remaining += f" (cause: {type(post_err.__cause__).__name__}: {post_err.__cause__})"
        raise RuntimeError(
            f"Workbook still failed to load in openpyxl after removing {len(removed)} definedName(s): {remaining}"
        ) from post_err

    removed_names = [entry["name"] for entry in removed]
    return {
        "status": "success",
        "input": str(src),
        "output": str(dst),
        "removed_count": len(removed),
        "removed_names": removed_names,
        "removed": removed,
    }


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Repair an .xlsx workbook by stripping definedNames that break openpyxl loading."
    )
    parser.add_argument("file", help="Path to the input .xlsx workbook")
    parser.add_argument("--out", default=None, help="Optional output path (defaults to <stem>.repaired.xlsx)")
    args = parser.parse_args(argv)

    try:
        report = repair_workbook(args.file, args.out)
    except Exception as exc:
        print(json.dumps({"status": "error", "error": str(exc)}, ensure_ascii=False))
        return 1

    print(json.dumps(report, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
