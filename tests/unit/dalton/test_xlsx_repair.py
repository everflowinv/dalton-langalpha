from __future__ import annotations

import importlib.util
import json
import zipfile
from pathlib import Path

import openpyxl
from openpyxl.workbook.defined_name import DefinedName
import pytest

REPO_ROOT = Path(__file__).resolve().parents[3]
REPAIR_SCRIPT = (
    REPO_ROOT
    / "plugins"
    / "langalpha_deliverables"
    / "skills"
    / "xlsx"
    / "scripts"
    / "repair.py"
)


def _load_repair_module():
    spec = importlib.util.spec_from_file_location("xlsx_repair", REPAIR_SCRIPT)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def _inject_xml_defined_names(xlsx_path: Path, extra_xml: str) -> None:
    """Inject raw `<definedName>` XML entries into `xl/workbook.xml`."""
    with zipfile.ZipFile(xlsx_path, "r") as zin:
        items = [(info, zin.read(info.filename)) for info in zin.infolist()]

    with zipfile.ZipFile(xlsx_path, "w", compression=zipfile.ZIP_DEFLATED) as zout:
        for info, data in items:
            if info.filename == "xl/workbook.xml":
                xml_text = data.decode("utf-8")
                assert "</definedNames>" in xml_text
                xml_text = xml_text.replace(
                    "</definedNames>", f"{extra_xml}</definedNames>", 1
                )
                data = xml_text.encode("utf-8")
            zout.writestr(info, data)


def test_repair_no_op_when_workbook_already_loads(
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    valid_xlsx = tmp_path / "clean_model.xlsx"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Model"
    ws["A1"] = "Revenue"
    ws["B1"] = 1250
    wb.defined_names.add(DefinedName(name="Base", attr_text="'Model'!$B$1"))
    wb.defined_names.add(DefinedName(name="Alias", attr_text="Base"))
    wb.defined_names.add(
        DefinedName(name="SafeFormula", attr_text='IFERROR(Model!$A$1,"#REF!")')
    )
    wb.save(valid_xlsx)
    original_bytes = valid_xlsx.read_bytes()

    repair_mod = _load_repair_module()
    exit_code = repair_mod.main([str(valid_xlsx)])
    assert exit_code == 0

    captured = capsys.readouterr().out.strip()
    assert "no repair needed" in captured
    assert valid_xlsx.read_bytes() == original_bytes
    assert not (tmp_path / "clean_model.repaired.xlsx").exists()


def test_repair_strips_corrupt_print_titles_and_preserves_aliases_and_formula_strings(
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    corrupt_xlsx = tmp_path / "google_model.xlsx"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Model"
    ws["A1"] = "Revenue"
    ws["B1"] = 1250
    wb.defined_names.add(DefinedName(name="Base", attr_text="'Model'!$B$1"))
    wb.defined_names.add(DefinedName(name="Alias", attr_text="Base"))
    wb.defined_names.add(
        DefinedName(name="SafeFormula", attr_text='IFERROR(Model!$A$1,"#REF!")')
    )
    wb.save(corrupt_xlsx)

    _inject_xml_defined_names(
        corrupt_xlsx,
        (
            '<definedName name="_xlnm.Print_Titles" localSheetId="0">"Formula removed..."</definedName>'
            '<definedName name="OutOfBoundsLocal" localSheetId="9">\'Model\'!$A$1</definedName>'
            # Parses in openpyxl (empty result), so it does not block loading and must be kept.
            '<definedName name="_xlnm.Print_Area" localSheetId="0">!!!</definedName>'
        ),
    )
    original_bytes = corrupt_xlsx.read_bytes()

    # Confirm openpyxl fails to load the corrupted workbook before repair.
    with pytest.raises(ValueError, match="could not assign names"):
        openpyxl.load_workbook(corrupt_xlsx)

    repair_mod = _load_repair_module()
    exit_code = repair_mod.main([str(corrupt_xlsx)])
    assert exit_code == 0

    captured = capsys.readouterr().out.strip()
    report = json.loads(captured)
    assert report["status"] == "success"
    assert report["removed_names"] == ["_xlnm.Print_Titles", "OutOfBoundsLocal"]

    # Original file must remain untouched.
    assert corrupt_xlsx.read_bytes() == original_bytes

    repaired_xlsx = tmp_path / "google_model.repaired.xlsx"
    assert repaired_xlsx.is_file()
    assert Path(report["output"]) == repaired_xlsx.resolve()

    # Repaired workbook loads cleanly in openpyxl and preserves Base, Alias=Base, and IFERROR(...,"#REF!").
    repaired_wb = openpyxl.load_workbook(repaired_xlsx)
    assert repaired_wb["Model"]["A1"].value == "Revenue"
    assert repaired_wb["Model"]["B1"].value == 1250
    assert "Base" in repaired_wb.defined_names
    assert "Alias" in repaired_wb.defined_names
    assert repaired_wb.defined_names["Alias"].attr_text == "Base"
    assert "SafeFormula" in repaired_wb.defined_names
    assert repaired_wb.defined_names["SafeFormula"].attr_text == 'IFERROR(Model!$A$1,"#REF!")'
    assert "OutOfBoundsLocal" not in repaired_wb.defined_names
